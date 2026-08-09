"""Import layer for normalized ICS topology ingestion.

This module keeps the Bayesian engine independent from raw engineering files.
Each supported file format is converted into a normalized topology dict with
assets and relationships before downstream processing.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

import networkx as nx
import yaml

try:
    import vsdx
except ImportError:  # pragma: no cover
    vsdx = None

from backend.topology import (
    DEFAULT_REL_TYPE,
    VALID_KINDS,
    normalize_asset,
    normalize_relationship,
    parse_generic_json,
    validate_graph,
)

_SUPPORTED_EXTENSIONS = {
    ".json",
    ".yaml",
    ".yml",
    ".csv",
    ".xlsx",
    ".vsdx",
    ".vdx",
    ".graphml",
    ".xml",
    ".aml",
}

_VSDX_SUPPORT_MESSAGE = (
    "Legacy binary Visio .vsd files cannot be parsed natively. "
    "Convert the file to a supported format first:\n"
    "  1. In Microsoft Visio: File > Save As > .vsdx\n"
    "  2. With LibreOffice: soffice --headless --convert-to vsdx your_file.vsd\n"
    "  3. Export the diagram to JSON, CSV, or GraphML from Visio and upload that.\n"
    "Supported formats: .json, .yaml, .yml, .csv, .xlsx, .graphml, .xml, .aml, .vsdx, .vdx"
)

_ASSET_HEADERS = {
    "id",
    "asset",
    "name",
    "label",
    "asset_id",
    "assetname",
    "node",
    "node_id",
}

_REL_HEADERS = {
    "source",
    "from",
    "target",
    "to",
    "dst",
    "destination",
    "destination_asset",
}

_ASSET_FIELD_MAP = {
    "id": "id",
    "asset": "id",
    "name": "name",
    "label": "label",
    "asset_id": "asset_id",
    "assetname": "assetName",
    "node": "node",
    "node_id": "node_id",
    "type": "type",
    "kind": "kind",
    "vendor": "vendor",
    "model": "model",
    "zone": "zone",
    "network": "network",
    "ip": "ip",
    "protocols": "protocols",
    "metadata": "metadata",
    # Cybersecurity attributes (must reach normalize_asset for strict
    # validation; previously these CSV columns were silently dropped).
    "cvss_type": "cvss_type",
    "cvss": "cvss_type",
    "exposed": "exposed",
    "patched": "patched",
    "consequence_severity": "consequence_severity",
    "consequence": "consequence_severity",
    "impact": "consequence_severity",
    "scope": "scope",
    "p_base_override": "p_base_override",
    "awareness": "awareness",
    "privilege": "privilege",
    "role": "role",
}

_REL_FIELD_MAP = {
    "source": "source",
    "from": "source",
    "target": "target",
    "to": "target",
    "type": "type",
    "rel_type": "type",
    "relationship_type": "type",
    "firewalled": "firewalled",
    "protected": "firewalled",
    "protocol": "protocol",
    "trust": "trust",
    "trust_level": "trust_level",
    "mitre": "mitre",
    "mitre_technique": "mitre_technique",
    "metadata": "metadata",
}


def load_topology(path: str | Path | dict) -> dict[str, Any]:
    """Load and normalize a topology from a path or inline dict.

    Returns ``{"assets", "relationships", "warnings"}`` where ``warnings`` is
    a list of human-readable notes about records that were normalized or
    skipped (e.g. self-loops, duplicate edges, unidentifiable assets).
    """
    if isinstance(path, dict):
        return _normalize_topology(path, "inline topology")

    path = Path(path)
    raw = _load_file(path)
    return _normalize_topology(raw, str(path))


def load_topology_from_bytes(content: bytes, filename: str) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".json", ".yaml", ".yml"}:
        raw = _parse_json_yaml(content, suffix)
    elif suffix == ".csv":
        raw = _parse_csv_bytes(content)
    elif suffix == ".xlsx":
        raw = _parse_excel_bytes(content, filename)
    elif suffix == ".graphml":
        raw = _parse_graphml_bytes(content)
    elif suffix in {".xml", ".aml"}:
        raw = _parse_aml_bytes(content)
    elif suffix == ".vdx":
        raw = _parse_vdx_bytes(content)
    elif suffix == ".vsd":
        raise ValueError(_VSDX_SUPPORT_MESSAGE)
    elif suffix == ".vsdx":
        raw = _parse_vsdx_bytes(content, filename)
    else:
        raise ValueError(
            f"Unsupported topology format '{suffix}'. Supported formats: "
            ".json, .yaml, .yml, .csv, .xlsx, .graphml, .xml, .aml, .vsdx, .vdx"
        )
    return _normalize_topology(raw, filename)


def _check_archive_expansion_limit(content: bytes, filename: str) -> None:
    """Reject zip-based uploads (xlsx/vsdx) that would decompress too far.

    A small zip can contain gigabytes of uncompressed data (zip bomb).  The
    API-level size limit covers the compressed payload; this check bounds the
    *uncompressed* total before any parser reads the archive.
    """
    max_expansion = int(
        os.getenv("MAX_ARCHIVE_EXPANSION_MB", "200")
    ) * 1024 * 1024
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total = sum(info.file_size for info in archive.infolist())
            if total > max_expansion:
                raise ValueError(
                    f"{filename}: the archive expands to {total / (1024 * 1024):.1f} MB "
                    f"uncompressed, exceeding the {max_expansion // (1024 * 1024)} MB "
                    "expansion limit. This file is rejected as a potential archive bomb."
                )
    except zipfile.BadZipFile:
        # Not a zip at all: let the real parser produce the format error.
        return


def _load_file(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in _SUPPORTED_EXTENSIONS or suffix == ".vsd":
        # .vsd is routed through load_topology_from_bytes so users get the
        # actionable conversion guidance rather than a generic format error.
        with open(path, "rb") as handle:
            return load_topology_from_bytes(handle.read(), path.name)
    raise ValueError(
        f"Unsupported topology format '{suffix}'. Supported formats: "
        ".json, .yaml, .yml, .csv, .xlsx, .graphml, .xml, .aml, .vsdx, .vdx"
    )


def _parse_json_yaml(content: bytes, suffix: str) -> Any:
    text = content.decode("utf-8-sig")
    if suffix == ".json":
        return json.loads(text)
    return yaml.safe_load(text)


def _read_csv_rows(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    return [[cell.strip() for cell in row] for row in reader]


def _split_csv_groups(rows: list[list[str]]) -> list[list[list[str]]]:
    groups: list[list[list[str]]] = []
    current: list[list[str]] = []
    for row in rows:
        if not any(cell for cell in row):
            if current:
                groups.append(current)
                current = []
            continue
        current.append(row)
    if current:
        groups.append(current)
    return groups


def _parse_csv_group(rows: list[list[str]]) -> tuple[dict[str, dict], list[Any]]:
    if not rows:
        return {}, []

    header = [cell.lower() for cell in rows[0]]
    has_asset_header = any(field in header for field in _ASSET_HEADERS)
    has_rel_header = any(field in header for field in _REL_HEADERS)

    assets: dict[str, dict] = {}
    relationships: list[Any] = []
    for row in rows[1:]:
        if not any(cell for cell in row):
            continue
        if has_asset_header:
            asset = _parse_csv_row_to_asset(row, header)
            if asset and asset.get("id"):
                assets[asset["id"]] = asset
        if has_rel_header:
            rel = _parse_csv_row_to_relationship(row, header)
            if rel:
                relationships.append(rel)

    if not has_rel_header and not has_asset_header:
        for row in rows[1:]:
            if len(row) >= 2 and row[0] and row[1]:
                rel_type = row[2] if len(row) > 2 else "connects-to"
                metadata: dict[str, Any] = {}
                if len(row) > 3 and row[3]:
                    metadata["protocol"] = row[3]
                if len(row) > 4 and row[4]:
                    metadata["trust_level"] = row[4]
                relationships.append({"source": row[0], "target": row[1], "type": rel_type, "metadata": metadata})

    return assets, relationships


def _parse_csv_bytes(content: bytes) -> dict[str, Any]:
    rows = _read_csv_rows(content)
    if not rows:
        raise ValueError("CSV topology file is empty.")

    groups = _split_csv_groups(rows)
    assets: dict[str, dict] = {}
    relationships: list[Any] = []
    for group in groups:
        group_assets, group_rels = _parse_csv_group(group)
        assets.update(group_assets)
        relationships.extend(group_rels)

    if not assets and not relationships:
        raise ValueError("Could not detect assets or relationships in the provided CSV file.")

    return {"assets": assets, "relationships": relationships}


def _parse_csv_row_to_asset(row: list[str], header: list[str]) -> dict[str, Any] | None:
    if not header:
        if len(row) < 2:
            return None
        return {"id": row[0].strip(), "name": row[1].strip()}

    raw: dict[str, Any] = {}
    for index, field in enumerate(header):
        if index >= len(row):
            continue
        key = _ASSET_FIELD_MAP.get(field)
        if not key:
            continue
        raw[key] = row[index]
    return raw if raw.get("id") or raw.get("name") else None


def _parse_csv_row_to_relationship(row: list[str], header: list[str]) -> dict[str, Any] | None:
    if not header:
        if len(row) < 2:
            return None
        return {"source": row[0].strip(), "target": row[1].strip()}

    raw: dict[str, Any] = {}
    for index, field in enumerate(header):
        if index >= len(row):
            continue
        key = _REL_FIELD_MAP.get(field)
        if not key:
            continue
        raw[key] = row[index]
    return raw if raw.get("source") and raw.get("target") else None


def _parse_excel_bytes(content: bytes, filename: str = "topology.xlsx") -> dict[str, Any]:
    _check_archive_expansion_limit(content, filename)
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "openpyxl is required to parse .xlsx files. Install with: pip install openpyxl"
        ) from exc

    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    assets: dict[str, dict] = {}
    relationships: list[Any] = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([str(cell).strip() if cell is not None else "" for cell in row])
        if not rows:
            continue
        groups = _split_csv_groups(rows)
        for group in groups:
            group_assets, group_rels = _parse_csv_group(group)
            assets.update(group_assets)
            relationships.extend(group_rels)
    workbook.close()
    if not assets and not relationships:
        raise ValueError("Excel topology file is empty or contains unsupported data.")
    return {"assets": assets, "relationships": relationships}


def _parse_graphml_bytes(content: bytes) -> dict[str, Any]:
    graph = nx.read_graphml(io.BytesIO(content))
    assets: dict[str, dict] = {}
    relationships: list[Any] = []
    _PROMOTED_NODE_ATTRS = {
        "label", "name", "kind", "type", "vendor", "model", "zone", "ip",
        # Security attributes: promote so they reach normalize_asset's strict
        # validation instead of being silently buried in metadata.
        "cvss_type", "cvss", "exposed", "patched",
        "consequence_severity", "consequence", "scope", "awareness",
        "privilege", "role", "p_base_override", "vulnerabilities",
    }
    for node, data in graph.nodes(data=True):
        raw: dict[str, Any] = {
            "id": node,
            "name": data.get("label") or data.get("name") or node,
            "kind": data.get("kind") or data.get("type"),
            "vendor": data.get("vendor"),
            "model": data.get("model"),
            "zone": data.get("zone"),
            "ip": data.get("ip"),
        }
        # Promote security attributes so they reach normalize_asset's strict
        # validation instead of being silently buried in metadata.  Only keys
        # actually present in the file are set (a bare None would be rejected
        # by the strict validators).
        promoted: dict[str, Any] = {
            "cvss_type": data.get("cvss_type", data.get("cvss")),
            "exposed": data.get("exposed"),
            "patched": data.get("patched"),
            "consequence_severity": data.get(
                "consequence_severity", data.get("consequence")
            ),
            "scope": data.get("scope"),
            "awareness": data.get("awareness"),
            "privilege": data.get("privilege"),
            "role": data.get("role"),
            "p_base_override": data.get("p_base_override"),
            "vulnerabilities": data.get("vulnerabilities"),
        }
        for key, value in promoted.items():
            if value is not None:
                raw[key] = value
        raw["metadata"] = {
            k: v for k, v in data.items() if k not in _PROMOTED_NODE_ATTRS
        }
        assets[node] = raw

    for source, target, data in graph.edges(data=True):
        raw = {
            "source": source,
            "target": target,
            "type": data.get("type") or data.get("relationship_type") or DEFAULT_REL_TYPE,
            "firewalled": data.get("firewalled", False),
            "protocol": data.get("protocol"),
            "trust_level": data.get("trust_level") or data.get("trust"),
            "mitre_technique": data.get("mitre_technique") or data.get("mitre"),
            "metadata": {k: v for k, v in data.items() if k not in {"type", "relationship_type", "firewalled", "protocol", "trust", "trust_level", "mitre_technique", "mitre"}},
        }
        relationships.append(raw)

    return {"assets": assets, "relationships": relationships}


def _parse_generic_xml(content: bytes) -> dict[str, Any]:
    text = content.decode("utf-8-sig", errors="replace")
    root = ET.fromstring(text)

    def strip_ns(tag: str) -> str:
        return tag[tag.find("}") + 1 :] if "}" in tag else tag

    def node_to_raw(element: ET.Element) -> dict[str, Any]:
        raw: dict[str, Any] = {}
        raw.update({strip_ns(k).lower(): v for k, v in element.attrib.items()})
        for child in element:
            child_tag = strip_ns(child.tag).lower()
            if child_tag and child.text:
                raw[child_tag] = child.text.strip()
        if "id" not in raw and "name" in raw:
            raw["id"] = raw["name"]
        return raw

    assets: dict[str, dict] = {}
    relationships: list[Any] = []
    asset_nodes: list[ET.Element] = []
    rel_nodes: list[ET.Element] = []

    root_tag = strip_ns(root.tag).lower()
    if root_tag in {"additionalmarkuplanguage", "automationml"}:
        return _parse_aml_bytes(content)

    for child in root:
        tag = strip_ns(child.tag).lower()
        if tag in {"assets", "nodes", "devices", "components", "items", "system_units", "internalelements", "internalelement"}:
            asset_nodes.extend(list(child))
        elif tag in {"relationships", "edges", "links", "connections", "paths", "connectionelements"}:
            rel_nodes.extend(list(child))
        elif tag in {"asset", "node", "device", "component", "item", "internalelement"}:
            asset_nodes.append(child)
        elif tag in {"relationship", "edge", "link", "connection"}:
            rel_nodes.append(child)

    if not asset_nodes:
        for element in root.iter():
            tag = strip_ns(element.tag).lower()
            if tag in {"asset", "node", "device", "component", "item", "internalelement"}:
                asset_nodes.append(element)
            elif tag in {"relationship", "edge", "link", "connection"}:
                rel_nodes.append(element)

    for element in asset_nodes:
        raw = node_to_raw(element)
        if raw.get("id"):
            assets[raw["id"]] = raw

    for element in rel_nodes:
        raw = node_to_raw(element)
        if raw.get("source") and raw.get("target"):
            relationships.append(raw)

    if not assets:
        raise ValueError("Could not parse any assets from the provided XML document.")
    return {"assets": assets, "relationships": relationships}


def _parse_aml_bytes(content: bytes) -> dict[str, Any]:
    text = content.decode("utf-8-sig", errors="replace")
    root = ET.fromstring(text)
    assets: dict[str, dict] = {}
    relationships: list[Any] = []

    def strip_ns(tag: str) -> str:
        return tag[tag.find("}") + 1 :] if "}" in tag else tag

    for element in root.iter():
        tag = strip_ns(element.tag)
        if tag == "InternalElement":
            # IEC 62714 semantics: `ID` is the canonical machine identifier,
            # `Name` the human-readable label, and `Connection` elements
            # reference IDs.  Prefer ID when present; fall back to Name so
            # ID-less documents (legacy AML exports) still parse.
            asset_id = element.get("ID") or element.get("Name") or element.get("NameLong")
            if not asset_id:
                continue
            asset_id = _normalize_text(asset_id)
            display_name = _normalize_text(element.get("Name") or asset_id)
            raw: dict[str, Any] = {"id": asset_id, "name": display_name}
            for child in element:
                child_tag = strip_ns(child.tag)
                if child_tag == "Attribute":
                    attr_name = child.get("Name")
                    attr_value = child.text
                    if attr_name and attr_value:
                        raw[attr_name.lower()] = _normalize_text(attr_value)
            assets[asset_id] = raw
        if tag in {"Connection", "InternalLink", "ConnectionElement"}:
            source = None
            target = None
            rel_type = DEFAULT_REL_TYPE
            metadata: dict[str, Any] = {}
            for child in element:
                child_tag = strip_ns(child.tag)
                if child_tag in {"Source", "From"}:
                    source = _normalize_text(child.text)
                elif child_tag in {"Target", "To"}:
                    target = _normalize_text(child.text)
                elif child_tag in {"Role", "Type"}:
                    rel_type = _normalize_text(child.text) or rel_type
                elif child_tag == "Protocol":
                    metadata["protocol"] = _normalize_text(child.text)
                elif child_tag in {"Trust_level", "Trust"}:
                    metadata["trust_level"] = _normalize_text(child.text)
            if source and target:
                relationships.append({"source": source, "target": target, "type": rel_type, "metadata": metadata})

    if not assets:
        try:
            return _parse_generic_xml(content)
        except ValueError:
            raise ValueError("Could not parse any assets from the AutomationML/XML document.")
    return {"assets": assets, "relationships": relationships}


def _parse_vdx_bytes(content: bytes) -> dict[str, Any]:
    """Parse Visio 2003-2010 XML (.vdx) format.

    .vdx is an XML-based format (not ZIP-based like .vsdx).  We extract
    shape text that follows the 'asset,...' or 'relationship,...' convention.
    """
    text = content.decode("utf-8-sig", errors="replace")
    root = ET.fromstring(text)

    def strip_ns(tag: str) -> str:
        return tag[tag.find("}") + 1 :] if "}" in tag else tag

    assets: dict[str, dict] = {}
    relationships: list[dict] = []

    # Find all Shape elements across all pages
    for shape in root.iter():
        tag = strip_ns(shape.tag)
        if tag != "Shape":
            continue

        # Extract text from the shape
        shape_text = ""
        for text_elem in shape.iter():
            if strip_ns(text_elem.tag) in {"Text", "Cp"}:
                if text_elem.text:
                    shape_text += text_elem.text

        shape_text = shape_text.strip()
        if not shape_text:
            continue

        # Check for asset/relationship markers
        if shape_text.lower().startswith("relationship,"):
            parts = [part.strip() for part in shape_text.split(",")]
            rel = {
                "source": parts[1] if len(parts) > 1 else None,
                "target": parts[2] if len(parts) > 2 else None,
                "type": parts[3] if len(parts) > 3 else DEFAULT_REL_TYPE,
                "firewalled": parts[4].lower() in {"1", "true", "yes"} if len(parts) > 4 else False,
                "protocol": parts[5] if len(parts) > 5 else None,
                "trust_level": parts[6] if len(parts) > 6 else None,
                "mitre_technique": parts[7] if len(parts) > 7 else None,
            }
            if rel["source"] and rel["target"]:
                relationships.append(rel)
            continue

        if shape_text.lower().startswith("asset,"):
            parts = [part.strip() for part in shape_text.split(",")]
            raw = {"id": parts[1] if len(parts) > 1 else None, "name": parts[1] if len(parts) > 1 else None}
            asset_kind = parts[2].lower() if len(parts) > 2 else ""
            if asset_kind in VALID_KINDS:
                raw["kind"] = asset_kind
            elif asset_kind:
                raw["type"] = asset_kind
            # Fields after the kind are interpreted per asset kind so the
            # Visio annotation convention (see tests/generate_vsdx.py)
            # round-trips the same security attributes as JSON/YAML/CSV.
            if asset_kind == "human":
                if len(parts) > 3:
                    raw["role"] = parts[3]
                if len(parts) > 4:
                    raw["awareness"] = parts[4]
                if len(parts) > 5:
                    raw["privilege"] = parts[5]
                if len(parts) > 6:
                    raw["consequence_severity"] = parts[6]
            elif asset_kind == "physical":
                if len(parts) > 3:
                    raw["p_base_override"] = parts[3]
                if len(parts) > 4:
                    raw["consequence_severity"] = parts[4]
            else:  # device (or legacy shapes without a valid kind token)
                if len(parts) > 3:
                    raw["cvss_type"] = parts[3]
                if len(parts) > 4:
                    raw["exposed"] = parts[4]
                if len(parts) > 5:
                    raw["patched"] = parts[5]
                if len(parts) > 6:
                    raw["consequence_severity"] = parts[6]
            asset_id = raw.get("id")
            if asset_id:
                assets[asset_id] = raw
            continue

        # Fallback: try shape custom properties (Prop elements)
        props = {}
        for prop in shape.iter():
            if strip_ns(prop.tag) == "Prop":
                label = prop.get("Label") or prop.get("N")
                val_elem = prop.find(".//{*}Value")
                if label and val_elem is not None and val_elem.text:
                    props[label] = val_elem.text.strip()

        if props.get("ID") or props.get("id") or props.get("Name"):
            raw = {
                "id": _normalize_text(props.get("ID") or props.get("id") or props.get("Name")),
                "name": _normalize_text(props.get("Name") or props.get("ID")),
            }
            if kind := props.get("Kind") or props.get("Type"):
                raw["kind"] = kind
            if vendor := props.get("Vendor"):
                raw["vendor"] = vendor
            if model := props.get("Model"):
                raw["model"] = model
            if raw["id"]:
                assets[raw["id"]] = raw

    if not assets:
        raise ValueError(
            "No asset shapes found in .vdx file. Ensure each asset shape contains readable text "
            "like 'asset,,,...' or add shape custom properties (ID/Name)."
        )
    return {"assets": assets, "relationships": relationships}


def _parse_vsdx_bytes(content: bytes, filename: str) -> dict[str, Any]:
    """Parse Visio .vsdx (Open XML) files."""
    _check_archive_expansion_limit(content, filename)
    if vsdx is None:
        raise ImportError(
            "vsdx library is required to parse .vsdx files. Install with: pip install vsdx"
        )

    with tempfile.NamedTemporaryFile(suffix=".vsdx", delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name

    try:
        try:
            doc = vsdx.VisioFile(tmp_path)
        except (TypeError, zipfile.BadZipFile) as exc:
            raise ValueError(
                "Could not parse the Visio file. Ensure the upload is a valid modern .vsdx package "
                "with readable shape text or properties."
            ) from exc

        assets: dict[str, dict] = {}
        relationships: list[Any] = []

        for page in doc.pages:
            for shape in page.shapes:
                text = ""
                for attr in ("text", "shape_text", "name"):
                    if hasattr(shape, attr):
                        text = getattr(shape, attr) or text
                text = _normalize_text(text)
                props = {}
                for attr in ("properties", "shape_data", "data"):
                    if hasattr(shape, attr):
                        raw_props = getattr(shape, attr)
                        if isinstance(raw_props, dict):
                            props.update(raw_props)

                if text:
                    if text.lower().startswith("relationship,"):
                        parts = [part.strip() for part in text.split(",")]
                        rel = {
                            "source": parts[1] if len(parts) > 1 else None,
                            "target": parts[2] if len(parts) > 2 else None,
                            "type": parts[3] if len(parts) > 3 else DEFAULT_REL_TYPE,
                            "firewalled": parts[4].lower() in {"1", "true", "yes"} if len(parts) > 4 else False,
                            "protocol": parts[5] if len(parts) > 5 else None,
                            "trust_level": parts[6] if len(parts) > 6 else None,
                            "mitre_technique": parts[7] if len(parts) > 7 else None,
                        }
                        if rel["source"] and rel["target"]:
                            relationships.append(rel)
                        continue

                    if text.lower().startswith("asset,"):
                        parts = [part.strip() for part in text.split(",")]
                        raw = {"id": parts[1] if len(parts) > 1 else None, "name": parts[1] if len(parts) > 1 else None}
                        asset_kind = parts[2].lower() if len(parts) > 2 else ""
                        if asset_kind in VALID_KINDS:
                            raw["kind"] = asset_kind
                        elif asset_kind:
                            raw["type"] = asset_kind
                        if asset_kind == "human":
                            if len(parts) > 3:
                                raw["role"] = parts[3]
                            if len(parts) > 4:
                                raw["awareness"] = parts[4]
                            if len(parts) > 5:
                                raw["privilege"] = parts[5]
                            if len(parts) > 6:
                                raw["consequence_severity"] = parts[6]
                        elif asset_kind == "physical":
                            if len(parts) > 3:
                                raw["p_base_override"] = parts[3]
                            if len(parts) > 4:
                                raw["consequence_severity"] = parts[4]
                        else:  # device (or legacy shapes without a valid kind token)
                            if len(parts) > 3:
                                raw["cvss_type"] = parts[3]
                            if len(parts) > 4:
                                raw["exposed"] = parts[4]
                            if len(parts) > 5:
                                raw["patched"] = parts[5]
                            if len(parts) > 6:
                                raw["consequence_severity"] = parts[6]
                        asset_id = raw.get("id")
                        if asset_id:
                            assets[asset_id] = raw
                        continue

                if props.get("ID") or props.get("id") or props.get("Name"):
                    raw = {
                        "id": _normalize_text(props.get("ID") or props.get("id") or props.get("Name")),
                        "name": _normalize_text(props.get("Name") or props.get("ID")),
                    }
                    if kind := props.get("Kind") or props.get("Type"):
                        raw["kind"] = kind
                    if vendor := props.get("Vendor"):
                        raw["vendor"] = vendor
                    if model := props.get("Model"):
                        raw["model"] = model
                    if raw["id"]:
                        assets[raw["id"]] = raw
                    continue

        if not assets:
            # Fallback: attempt to scan the VSDX package XML files for labeled
            # text like 'asset,,...' or 'relationship,,,...'. This
            # helps when the `vsdx` library cannot expose shape text but the
            # raw XML still contains our authoring markers.
            try:
                with zipfile.ZipFile(tmp_path) as zf:
                    for name in zf.namelist():
                        if not name.lower().endswith('.xml'):
                            continue
                        try:
                            xml_text = zf.read(name).decode('utf-8', errors='ignore')
                        except Exception:
                            continue
                        for m in re.finditer(r"(asset|relationship),([^<\r\n]+)", xml_text, flags=re.IGNORECASE):
                            parts = [p.strip() for p in m.group(2).split(',')]
                            if m.group(1).lower() == 'asset':
                                aid = parts[0] if parts else None
                                if aid:
                                    raw_asset = {"id": aid, "name": aid}
                                    if len(parts) > 1:
                                        raw_asset["type"] = parts[1]
                                    assets[aid] = raw_asset
                            else:  # relationship
                                src = parts[0] if len(parts) > 0 else None
                                tgt = parts[1] if len(parts) > 1 else None
                                if src and tgt:
                                    rel_type = parts[2] if len(parts) > 2 else DEFAULT_REL_TYPE
                                    firewalled = parts[3].lower() in {"1", "true", "yes"} if len(parts) > 3 else False
                                    relationships.append({
                                        "source": src,
                                        "target": tgt,
                                        "type": rel_type,
                                        "firewalled": firewalled,
                                    })
            except Exception:
                pass

        if not assets:
            raise ValueError(
                "No asset shapes found in .vsdx file. Ensure each asset shape contains readable text like 'asset,,,...' or add shape custom properties (ID/Name)."
                " If you cannot modify the Visio source, export the diagram to GraphML, JSON, or CSV from Visio and re-upload."
            )
        return {"assets": assets, "relationships": relationships}
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_topology(raw: Any, source_label: str) -> dict[str, Any]:
    if not isinstance(raw, dict):
        raise ValueError(f"{source_label}: topology payload must be an object.")

    warnings: list[str] = []

    if "assets" in raw and "relationships" in raw:
        assets_raw = raw["assets"]
        relationships_raw = raw["relationships"]
    else:
        inferred = parse_generic_json(raw, warnings)
        assets_raw = inferred["assets"]
        relationships_raw = inferred["relationships"]

    assets: dict[str, dict] = {}
    if isinstance(assets_raw, dict):
        for node_id, attrs in assets_raw.items():
            if not isinstance(attrs, dict):
                warnings.append(
                    f"{source_label}: asset '{node_id}' was skipped because its "
                    "attributes are not an object."
                )
                continue
            normalized = normalize_asset({**attrs, "id": node_id})
            if normalized:
                assets[normalized["id"]] = normalized
            else:
                warnings.append(
                    f"{source_label}: asset '{node_id}' was skipped because it has "
                    "no usable identifier."
                )
    elif isinstance(assets_raw, list):
        for item in assets_raw:
            if not isinstance(item, dict):
                warnings.append(
                    f"{source_label}: an asset record was skipped because it is "
                    "not an object."
                )
                continue
            normalized = normalize_asset(item)
            if normalized:
                assets[normalized["id"]] = normalized
            else:
                warnings.append(
                    f"{source_label}: asset record {item!r} was skipped because it "
                    "has no usable identifier."
                )
    elif assets_raw:
        raise ValueError(
            f"{source_label}: 'assets' must be an object or a list of objects."
        )

    relationships: list[tuple] = []
    if isinstance(relationships_raw, list):
        for item in relationships_raw:
            rel_normalized = normalize_relationship(item)
            if rel_normalized:
                relationships.append(rel_normalized)
            else:
                warnings.append(
                    f"{source_label}: relationship {item!r} was skipped because it "
                    "is malformed (missing source or target)."
                )
    elif relationships_raw:
        raise ValueError(
            f"{source_label}: 'relationships' must be a list."
        )

    relationships = validate_graph(assets, relationships, source_label, warnings)
    return {"assets": assets, "relationships": relationships, "warnings": warnings}